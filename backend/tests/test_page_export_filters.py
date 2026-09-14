"""Выгрузка страницы в Excel совпадает с тем, что человек видел на экране.

Два независимых дефекта, закрытые вместе (аудит 14.09):

1. Фильтры страницы (период, строка, ветка лестницы) до выгрузки не доходили
   вовсе: маршрут их не принимал. Руководитель отфильтровывал страницу на одно
   отделение за июль, выгружал Excel и нёс на совещание цифры по всей форме за
   последний отчёт — без единой оговорки в файле.

2. `compute_widget_data` вызывался без `user`, а набор разрешённых строк
   (`allowed_rows_for_dataset`) при `user=None` означает «без ограничения».
   То есть выгрузка отдавала ВСЕ строки формы тому, кому разрешена одна.
   `skip_acl` должен пропускать только повторную проверку ДАШБОРДА.

Читаем не код ответа, а сам файл: дефект был именно в содержимом."""
import io

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.asyncio(loop_scope="session")

from conftest import db, purge_dashboard

ALL_ROWS = {"Паспорт", "ИНН", "СНИЛС"}


def _cells(blob: bytes) -> set:
    """Все непустые значения всех листов книги — как строки."""
    wb = load_workbook(io.BytesIO(blob))
    out = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is not None and str(v).strip():
                    out.add(str(v).strip())
    return out


async def _page_with_table(client, headers, name: str):
    did = (await client.post("/dashboards", headers=headers, json={"name": name})).json()["id"]
    pid = (await client.post(f"/dashboards/{did}/pages", headers=headers, json={"name": "P"})).json()["id"]
    await client.post(f"/dashboard-pages/{pid}/widgets", headers=headers,
                      json={"name": "Таблица", "widget_type": "table", "config": {"dataset_code": "t_ds"}})
    return did, pid


async def test_row_filter_reaches_the_file(client, admin_headers, seed_dataset):
    """Фильтр строки сужает файл и назван в «Содержании»."""
    did, pid = await _page_with_table(client, admin_headers, "ztest_exp_filter")
    try:
        full = _cells((await client.get(f"/dashboard-pages/{pid}/export.xlsx",
                                        headers=admin_headers)).content)
        assert ALL_ROWS <= full, "без фильтра в файле должны быть все строки"

        one = _cells((await client.get(f"/dashboard-pages/{pid}/export.xlsx?row=Паспорт",
                                       headers=admin_headers)).content)
        assert "Паспорт" in one
        assert not ({"ИНН", "СНИЛС"} & one), "фильтр строки не дошёл до файла"
        # Оговорка о фильтре обязана быть в файле: иначе отфильтрованную
        # выгрузку не отличить от полной.
        assert "Строка" in one and "Фильтры, действовавшие при выгрузке" in one
    finally:
        await purge_dashboard(did)


async def test_period_filter_reaches_the_file(client, admin_headers, seed_dataset):
    """Период, в котором отчётов нет, не может дать в файле цифры последнего."""
    did, pid = await _page_with_table(client, admin_headers, "ztest_exp_period")
    try:
        blob = (await client.get(
            f"/dashboard-pages/{pid}/export.xlsx?from=2000-01-01&to=2000-12-31",
            headers=admin_headers)).content
        cells = _cells(blob)
        assert not (ALL_ROWS & cells), "за период без отчётов в файл попали чужие данные"
        assert "Период" in cells, "период не назван в «Содержании»"
    finally:
        await purge_dashboard(did)


async def test_export_respects_row_level_rls(client, admin_headers, viewer, seed_dataset):
    """Главное: выгрузка не обходит ПОСТРОЧНЫЕ права.

    Зрителю разрешена одна строка — столько же должно быть в файле.
    До правки в файл уезжали все три."""
    async with db.acquire() as conn:
        obj = str(await conn.fetchval("select id from objects where name='t_obj'"))
        org = await conn.fetchval("select organization_id from objects where id=$1::uuid", obj)
        dep = str(await conn.fetchval(
            "insert into departments(organization_id,name) values($1,'ztest_dep_exp') returning id", org))
        await conn.execute("update users set department_id=$1::uuid where id=$2::uuid", dep, viewer["id"])

    did, pid = await _page_with_table(client, admin_headers, "ztest_exp_rls")
    try:
        await client.post(f"/dashboards/{did}/grants", headers=admin_headers,
                          json={"grantee_type": "user", "user_id": viewer["id"]})
        await client.post(f"/dashboards/{did}/publish", headers=admin_headers)
        r = await client.put(f"/objects/{obj}/row-acl/{dep}", headers=admin_headers,
                             json={"row_labels": ["Паспорт"]})
        assert r.status_code == 200, r.text

        mine = _cells((await client.get(f"/dashboard-pages/{pid}/export.xlsx",
                                        headers=viewer["headers"])).content)
        assert "Паспорт" in mine, "разрешённая строка пропала из файла"
        assert not ({"ИНН", "СНИЛС"} & mine), "выгрузка обошла построчные права"

        # Привилегированному по-прежнему видно всё — правило не сломано в другую сторону
        boss = _cells((await client.get(f"/dashboard-pages/{pid}/export.xlsx",
                                        headers=admin_headers)).content)
        assert ALL_ROWS <= boss
    finally:
        async with db.acquire() as conn:
            await conn.execute("delete from data_row_acl where department_id=$1::uuid", dep)
            await conn.execute("update users set department_id=null where id=$1::uuid", viewer["id"])
            await conn.execute("delete from departments where id=$1::uuid", dep)
        await purge_dashboard(did)
