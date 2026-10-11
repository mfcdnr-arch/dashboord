"""Книга из нескольких форм: шаблон разметки — на ФОРМУ, а не на объект (миграция 061).

Недельный «Отчёт МФЦ для Минэкономразвития» — один объект и несколько форм на
разных листах, у каждой свой код набора. Пока шаблон был один на объект,
выпуск одного листа затирал разметку соседнего, и следующая книга узнавалась
по ОДНОМУ листу: недельная автозагрузка выпускала одну форму из четырёх.

Здесь проверяется весь путь: две формы одной книги размечаются один раз, и
следующая книга выпускает ОБЕ сама; неполный выпуск книги назван; форму можно
перестать узнавать и вернуть; ступени, таблица, проверка качества и мастер
сборки понимают, что форм несколько и что одна из них — реестр.
"""
import io

import pytest

pytestmark = pytest.mark.asyncio(loop_scope="session")

from app import db
from app.modules.dashboards import _suggest, _widgetsources
from app.modules.ingestion import form_templates as ft
from app.modules.ingestion import quality

from test_layout_template import _upload_and_extract, obj  # noqa: F401

POK, ITOGI = "ztb_pok", "ztb_itogi"


def _book(week: int, itogi_rows=None) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Показатели"
    ws.append(["№ п/п", "Субъект", "Обращения", "Уведомления"])
    ws.append(["1", "2", "3", "4"])
    ws.append(["1", "ДНР", 1000 + week, 100 + week])
    ws2 = wb.create_sheet("Итоги недели")
    ws2.append(["№ п/п", "Наименование", "Характеристика", "Ответственный", "Особо значимое", "Баллы"])
    ws2.append(["1", "2", "3", "4", "5", "6"])
    for r in itogi_rows or [
        ("1", "Открыт офис в Макеевке", "приём с 10.10", "Иванов", "да", 5),
        ("2", "Обучение операторов", "40 человек", "Петров", "нет", 3),
        ("3", "Новая услуга ФНС", "регистрация ИП", "Сидоров", "да", 4),
        ("4", f"Совещание недели {week}", "по плану", "Иванов", "нет", 2),
    ]:
        ws2.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _body(table: dict, code: str, period: str, cols: list) -> dict:
    """cols: [(индекс столбца, код, имя, тип)]; первый — название строки."""
    return {
        "table_id": table["id"], "code": code, "name": table["sheet_or_page"],
        "reporting_period_start": period,
        "fields": [{"column_index": i, "field_code": c, "field_name": n, "data_type": t,
                    "is_row_label": k == 0} for k, (i, c, n, t) in enumerate(cols)],
        "layout": {"data_rect": table["data_rect"], "header_rows": 2,
                   "orientation": "columns", "skip_rows": []},
    }


POK_COLS = [(1, "zsubj", "Субъект", "text"), (2, "zobr", "Обращения", "number"),
            (3, "zuved", "Уведомления", "number")]
ITOGI_COLS = [(1, "zname", "Наименование", "text"), (2, "zchar", "Характеристика", "text"),
              (3, "zresp", "Ответственный", "text"), (4, "zsign", "Особо значимое", "text"),
              (5, "zball", "Баллы", "number")]


def _table(job: dict, sheet: str) -> dict:
    return next(t for t in job["tables"] if t["sheet_or_page"] == sheet)


async def _first_book(client, headers, obj, monkeypatch):  # noqa: F811
    job = await _upload_and_extract(client, headers, obj["folder_id"], _book(1), "2026-10-06", monkeypatch)
    for sheet, code, cols in (("Показатели", POK, POK_COLS), ("Итоги недели", ITOGI, ITOGI_COLS)):
        r = await client.post(f"/extraction-jobs/{job['job_id']}/release", headers=headers,
                              json=_body(_table(job, sheet), code, "2026-10-06", cols))
        assert r.status_code == 201, r.text
    return job


async def _active(code: str) -> list:
    async with db.acquire() as conn:
        rows = await conn.fetch(
            "select reporting_period_start p from dataset_releases where code=$1 and status<>'superseded' "
            "order by 1", code)
    return [r["p"].isoformat() for r in rows]


async def test_each_form_keeps_its_own_template_and_next_book_releases_all(
        client, admin_headers, obj, monkeypatch):  # noqa: F811
    await _first_book(client, admin_headers, obj, monkeypatch)
    async with db.acquire() as conn:
        tpls = await ft.for_object(conn, obj["object_id"])
    # Второй выпуск НЕ затёр разметку первого — у каждой формы свой шаблон.
    assert {(t["dataset_code"], t["sheet_name"]) for t in tpls} == {
        (POK, "Показатели"), (ITOGI, "Итоги недели")}

    job2 = await _upload_and_extract(client, admin_headers, obj["folder_id"], _book(2),
                                     "2026-10-13", monkeypatch)
    by_code = {e["dataset_code"]: e for e in job2["layout_templates"]}
    assert by_code[POK]["match"] == by_code[ITOGI]["match"] == "exact"
    assert by_code[POK]["table_id"] == _table(job2, "Показатели")["id"]
    assert by_code[ITOGI]["table_id"] == _table(job2, "Итоги недели")["id"]
    # Авто-выпуск выпустил ОБЕ формы, а не одну последнюю размеченную.
    assert await _active(POK) == ["2026-10-06", "2026-10-13"]
    assert await _active(ITOGI) == ["2026-10-06", "2026-10-13"]
    async with db.acquire() as conn:
        # Одно уведомление на файл, с обеими формами внутри.
        events = await conn.fetch(
            "select payload from notification_events where event_type='data.auto_released' "
            "and entity_id in (select id from dataset_releases where object_id=$1::uuid)", obj["object_id"])
    assert len(events) == 1
    import json
    payload = events[0]["payload"]
    payload = json.loads(payload) if isinstance(payload, str) else payload
    assert {f["code"] for f in payload["forms"]} == {POK, ITOGI}


async def test_partial_book_release_is_named(client, admin_headers, obj, monkeypatch):  # noqa: F811
    await _first_book(client, admin_headers, obj, monkeypatch)
    async with db.acquire() as conn:
        await conn.execute("update folders set auto_release=false where id=$1::uuid", obj["folder_id"])
    job = await _upload_and_extract(client, admin_headers, obj["folder_id"], _book(3),
                                    "2026-10-20", monkeypatch)
    tpl = next(e for e in job["layout_templates"] if e["dataset_code"] == POK)
    r = await client.post(f"/extraction-jobs/{job['job_id']}/release", headers=admin_headers,
                          json=_body(_table(job, "Показатели"), POK, "2026-10-20", POK_COLS)
                          | {"layout": tpl["layout"]})
    assert r.status_code == 201, r.text
    r = await client.get(f"/folders/{obj['folder_id']}/documents", headers=admin_headers)
    doc = next(d for d in r.json()["items"] if d["reporting_period_start"] == "2026-10-20")
    # «Данные выпущены» спрятало бы, что вторая форма книги не выпущена.
    assert doc["pipeline"] == "released_partial", doc
    assert "1 из 2" in doc["pipeline_hint"]


async def test_forget_and_restore_a_form(client, admin_headers, obj, monkeypatch):  # noqa: F811
    await _first_book(client, admin_headers, obj, monkeypatch)
    oid = obj["object_id"]
    r = await client.post(f"/objects/{oid}/layout-templates/forms/{ITOGI}/forget", headers=admin_headers)
    assert r.status_code == 200, r.text
    st = r.json()
    assert [t["dataset_code"] for t in st["templates"]] == [POK], "соседняя форма не тронута"
    hist = next(h for h in st["history"] if h["dataset_code"] == ITOGI)
    r = await client.post(f"/objects/{oid}/layout-templates/{hist['id']}/restore", headers=admin_headers)
    assert r.status_code == 200, r.text
    assert {t["dataset_code"] for t in r.json()["templates"]} == {POK, ITOGI}
    # Повторно «забыть» несуществующее — честный отказ.
    r = await client.post(f"/objects/{oid}/layout-templates/forms/zzz_none/forget", headers=admin_headers)
    assert r.status_code == 404


async def test_levels_table_title_registry_and_planner(client, admin_headers, obj, monkeypatch):  # noqa: F811
    await _first_book(client, admin_headers, obj, monkeypatch)
    oid = obj["object_id"]

    # Ступени: форм несколько — сервер не угадывает, какую имели в виду.
    r = await client.get(f"/objects/{oid}/levels", headers=admin_headers)
    assert r.status_code == 200 and r.json()["needs_form"] is True, r.text
    assert {f["dataset_code"] for f in r.json()["forms"]} == {POK, ITOGI}
    r = await client.get(f"/objects/{oid}/levels", params={"code": ITOGI}, headers=admin_headers)
    assert r.status_code == 200 and r.json()["dataset_code"] == ITOGI, r.text

    async with db.acquire() as conn:
        org = await conn.fetchval("select organization_id from objects where id=$1::uuid", oid)
        # Таблица подписывает строки графой-названием ЭТОГО выпуска, а не «Строка».
        table = await _widgetsources._dataset_table(conn, org, ITOGI)
        assert table["row_title"] == "Наименование"
        assert (await _widgetsources._dataset_table(conn, org, POK))["row_title"] == "Субъект"

        rel = await conn.fetchval("select id from dataset_releases where code=$1 and status<>'superseded'", ITOGI)
        assert await quality.release_is_registry(conn, rel) is True
        rel_pok = await conn.fetchval("select id from dataset_releases where code=$1 and status<>'superseded'", POK)
        assert await quality.release_is_registry(conn, rel_pok) is False

        plan = await _suggest.auto_build_plan(conn, org, oid, with_metrics=False)
    # Формы одной книги — не «разорванный ряд»: ложного предупреждения нет.
    assert not any("несколько наборов" in w for w in plan["warnings"]), plan["warnings"]
    itogi = [c for c in plan["candidates"] if c["dataset_code"] == ITOGI]
    # Реестр считается строками, а не суммой баллов.
    assert any(c["widget_type"] == "kpi" and c["config"].get("count") for c in itogi), itogi
    assert any(c["config"].get("group_by") == "zsign" for c in itogi), itogi
    assert not any(c["config"].get("value_field") == "zball" for c in itogi), itogi


def test_registry_rule():
    assert quality.is_registry(["text", "text", "text", "number"]) is True
    assert quality.is_registry(["text", "number", "number"]) is False
    assert quality.is_registry(["text", "number"]) is False, "две графы — не о чем судить"


def test_registry_does_not_flag_repeated_rows():
    """Вопрос реестра висит неделями с той же оценкой — это не «перенос цифр»."""
    cur = {("Вопрос А", "risk"): 15.0, ("Вопрос Б", "risk"): 4.0}
    names = {"risk": "Оценка риска"}
    plain = quality.compare_with_previous(cur, dict(cur), names, "29.09.2026")
    assert any(w["code"] == "same_as_previous" for w in plain)
    reg = quality.compare_with_previous(cur, dict(cur), names, "29.09.2026", registry=True)
    assert not any(w["code"] == "same_as_previous" for w in reg)


# ── Сверка со сводкой той же книги (11.10.2026) ──────────────────────────────

def test_summary_rule_matches_sheet_by_word_stems():
    """«Планов на неделю: 22» относится к листу «Планы на неделю» — падеж не мешает."""
    from app.modules.ingestion import mapping
    grid = [["Минэкономразвития ДНР · панель отчётности", ""], ["", ""],
            ["Показателей в форме", "0"], ["Среднее достижение плана", "—"],
            ["Ниже 70 % — зона контроля", "0"], ["Итогов недели", "22"], ["Планов на неделю", "22"]]
    sums = [("Сводка", label, n) for label, n in mapping.summary_counts(grid)]
    assert ("Сводка", "Ниже 70 % — зона контроля", 0) not in sums, "процент — не счётчик строк"
    bad = mapping.summary_sheet_mismatch("Планы на неделю", 21, sums)
    assert bad and bad["code"] == "summary_sheet_mismatch" and "22" in bad["message"] and "21" in bad["message"]
    assert mapping.summary_sheet_mismatch("Итоги недели", 22, sums) is None
    assert mapping.summary_sheet_mismatch("Показатели", 0, sums) is None
    # Лист без слов в имени (книга «лист = дата») не сверяется: сопоставлять не с чем.
    assert mapping.summary_sheet_mismatch("31.08", 5, sums) is None
    # Лист, о котором сводка молчит, не сверяется.
    assert mapping.summary_sheet_mismatch("Классификатор рисков", 30, sums) is None


async def test_summary_mismatch_stops_auto_release(client, admin_headers, obj, monkeypatch):  # noqa: F811
    """Книга противоречит своей сводке — форма уходит человеку, остальные выпускаются."""
    from openpyxl import load_workbook

    await _first_book(client, admin_headers, obj, monkeypatch)
    wb = load_workbook(io.BytesIO(_book(2)))
    ws = wb.create_sheet("Сводка", 0)
    ws.append(["Итогов недели", 5])          # на листе 4 строки
    ws.append(["Обращений всего", 1002])     # о листах книги не говорит — не сверяется
    buf = io.BytesIO()
    wb.save(buf)
    job = await _upload_and_extract(client, admin_headers, obj["folder_id"], buf.getvalue(),
                                    "2026-10-13", monkeypatch)
    assert await _active(POK) == ["2026-10-06", "2026-10-13"], "соседняя форма выпущена"
    assert await _active(ITOGI) == ["2026-10-06"], "форма, расходящаяся со сводкой, ждёт человека"

    # Перед кнопкой «Выпустить» модератор видит, в чём дело.
    tpl = next(e for e in job["layout_templates"] if e["dataset_code"] == ITOGI)
    r = await client.post(f"/extraction-jobs/{job['job_id']}/quality-check", headers=admin_headers,
                          json=_body(_table(job, "Итоги недели"), ITOGI, "2026-10-13", ITOGI_COLS)
                          | {"layout": tpl["layout"]})
    assert r.status_code == 200, r.text
    hit = [w for w in r.json()["warnings"] if w["code"] == "summary_sheet_mismatch"]
    assert hit and "Итогов недели: 5" in hit[0]["message"] and "строк 4" in hit[0]["message"], r.json()
