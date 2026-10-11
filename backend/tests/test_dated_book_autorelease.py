"""Книга «лист = дата» выпускается сама целиком; время «ч:мм:сс» — числом (11.10.2026).

До этого авто-выпуск брал из книги ОДИН лист — за отчётную дату файла, а
ежедневный отчёт, форма которого за год росла, не выпускался сам никогда.
Теперь выпускаются все новые и изменившиеся дни тем же кодом, что за кнопкой
«Выпустить по листам», плюс три правила авто-режима: область подгоняется под
каждый лист (форма растёт вниз и вправо), дни позже отчётной даты файла не
трогаются, а день с замечаниями качества остаётся человеку.

И отдельно — длительности: отчёт «Среднее время обслуживания / ожидания» пишет
время как «00:14:22», и до 11.10 оно ложилось текстом.
"""
import datetime as dt
import io
import json

import pytest

pytestmark = pytest.mark.asyncio(loop_scope="session")

from app import db
from app.modules.dashboards import _widgetsources
from app.modules.ingestion import analyze, parsers, service

from pipeline_helpers import _upload, folder, offline_queue  # noqa: F401

CODE = "ztest_dated_book"


def _book(sheets: dict, extra: dict | None = None) -> bytes:
    """{«дд.мм»: [(отделение, принято, выдано)]}; extra — {лист: графа, добавленная справа}."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title=title)
        add = (extra or {}).get(title)
        ws.append(["Отделение", "Принято", "Выдано"] + ([add] if add else []))
        for r in rows:
            ws.append(list(r) + ([7] if add else []))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _load(client, headers, folder_id, content, period, monkeypatch):
    up = await _upload(client, headers, folder_id, content, period, monkeypatch)
    await service.run_extraction(up["extraction_job_id"])
    return up["extraction_job_id"]


async def _active() -> dict:
    async with db.acquire() as conn:
        rows = await conn.fetch(
            "select r.reporting_period_start p, count(distinct v.row_index) n_rows, "
            "  count(distinct v.canonical_field_code) n_fields, "
            "  max(v.value_number) filter (where v.row_label='Отделение 1' and v.canonical_field_code like 'prin%') pr "
            "from dataset_releases r left join dataset_values v on v.dataset_release_id=r.id "
            "where r.code=$1 and r.status<>'superseded' group by 1 order by 1", CODE)
    return {r["p"].isoformat(): (r["n_rows"], r["n_fields"], r["pr"]) for r in rows}


def _rows(base: int, n: int = 3) -> list:
    return [(f"Отделение {i}", base + i, base // 2 + i) for i in range(1, n + 1)]


async def _first_book(client, admin_headers, folder, monkeypatch):  # noqa: F811
    """Первая книга размечается человеком — кнопкой «Выпустить по листам»."""
    job = await _load(client, admin_headers, folder["folder_id"],
                      _book({"01.07": _rows(100), "02.07": _rows(110)}), "2026-07-02", monkeypatch)
    r = await client.post(f"/extraction-jobs/{job}/release-by-sheet", headers=admin_headers, json={
        "code": CODE, "name": "Ежедневный", "year": 2026,
        "layout": {"data_rect": None, "header_rows": 1, "orientation": "columns", "skip_rows": []}})
    assert r.status_code == 201, r.text
    assert set((await _active())) == {"2026-07-01", "2026-07-02"}


async def test_next_book_releases_all_new_days_and_follows_form_growth(
        client, admin_headers, folder, monkeypatch, offline_queue):  # noqa: F811
    await _first_book(client, admin_headers, folder, monkeypatch)
    book = _book({
        "01.07": _rows(100), "02.07": _rows(110),          # без изменений
        "03.07": _rows(120, n=4),                           # новый день, новое отделение
        "04.07": _rows(130, n=4),                           # новый день, форма выросла вправо
        "06.07": _rows(150),                                # позже отчётной даты файла
    }, extra={"04.07": "Отказ"})
    await _load(client, admin_headers, folder["folder_id"], book, "2026-07-04", monkeypatch)
    act = await _active()
    assert set(act) == {"2026-07-01", "2026-07-02", "2026-07-03", "2026-07-04"}, \
        "выпущены оба новых дня, лист позже отчётной даты не тронут"
    assert act["2026-07-03"][0] == 4, "новое отделение внизу листа не обрезано областью шаблона"
    assert act["2026-07-04"][1] == act["2026-07-03"][1] + 1, "графа, добавленная справа, не потерялась"
    async with db.acquire() as conn:
        replaced = await conn.fetchval(
            "select count(*) from dataset_releases where code=$1 and status='superseded'", CODE)
        ev = await conn.fetchrow(
            "select payload from notification_events where event_type='data.auto_released' "
            "and payload->>'code'=$1 order by created_at desc limit 1", CODE)
    assert replaced == 0, "неизменившиеся дни не перевыпускаются"
    payload = json.loads(ev["payload"]) if isinstance(ev["payload"], str) else ev["payload"]
    assert payload["book"] and payload["days"] == 2 and payload["first"] == "2026-07-03"


async def test_day_with_quality_warning_is_left_to_a_person(
        client, admin_headers, folder, monkeypatch, offline_queue):  # noqa: F811
    await _first_book(client, admin_headers, folder, monkeypatch)
    # 03.07 повторяет 02.07 цифра в цифру — «перенесли без обновления?».
    book = _book({"01.07": _rows(100), "02.07": _rows(110), "03.07": _rows(110), "04.07": _rows(130)})
    await _load(client, admin_headers, folder["folder_id"], book, "2026-07-04", monkeypatch)
    act = await _active()
    assert "2026-07-03" not in act, "день с замечанием не выпущен молча"
    assert "2026-07-04" in act, "остальные дни книги выпущены"
    async with db.acquire() as conn:
        ev = await conn.fetchrow(
            "select payload from notification_events where event_type='data.auto_released' "
            "and payload->>'code'=$1 order by created_at desc limit 1", CODE)
    payload = json.loads(ev["payload"]) if isinstance(ev["payload"], str) else ev["payload"]
    assert payload["held_total"] == 1 and payload["held"][0]["period"] == "2026-07-03"


# ── Длительности ─────────────────────────────────────────────────────────────

def test_duration_parsing():
    assert analyze.parse_number("00:14:22") == pytest.approx(14 + 22 / 60)
    assert analyze.parse_number("01:27:16") == pytest.approx(87 + 16 / 60)
    assert analyze.parse_duration("08:00") is None, "без секунд — скорее время суток"
    assert analyze.is_duration_column(["00:14:22", "00:07:37", "-", ""])
    assert not analyze.is_duration_column(["14", "15", "00:01:00"])
    assert analyze.infer_type(["00:14:22", "00:07:37"])[0] == "number"


def test_excel_time_cells_become_durations():
    assert parsers._cell_to_str(dt.time(0, 14, 22, 500000)) == "00:14:22"
    assert parsers._cell_to_str(dt.timedelta(hours=1, minutes=2, seconds=3)) == "01:02:03"
    assert parsers._cell_to_str(dt.timedelta(days=1, minutes=1)) == "24:01:00"


async def test_duration_form_is_released_as_minutes(
        client, admin_headers, folder, monkeypatch, offline_queue, ids):  # noqa: F811
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "17.08.26"
    ws.append(["Название офиса", "Среднее время ожидания", "Всего выдано талонов"])
    ws.append(["Отделение 1", dt.time(0, 14, 22), 101])
    ws.append(["Отделение 2", dt.time(1, 27, 16), 227])
    buf = io.BytesIO()
    wb.save(buf)
    job = await _load(client, admin_headers, folder["folder_id"], buf.getvalue(), "2026-08-17", monkeypatch)
    r = await client.get(f"/extraction-jobs/{job}", headers=admin_headers)
    table = r.json()["tables"][0]
    r = await client.post(f"/extraction-jobs/{job}/layout-preview", headers=admin_headers, json={
        "table_id": table["id"], "data_rect": None, "header_rows": 1, "orientation": "columns", "skip_rows": []})
    cols = {c["field_name"]: c for c in r.json()["columns"]}
    wait = cols["Среднее время ожидания"]
    assert wait["data_type"] == "number" and wait["unit"] == "мин", wait
    assert "unit" not in cols["Всего выдано талонов"]

    fields = [{"column_index": c["column_index"], "field_code": c["field_code"], "field_name": c["field_name"],
               "data_type": c["data_type"], "is_row_label": c["column_index"] == 0,
               **({"unit": c["unit"]} if c.get("unit") else {})} for c in r.json()["columns"]]
    r = await client.post(f"/extraction-jobs/{job}/release", headers=admin_headers, json={
        "table_id": table["id"], "code": "ztest_wait", "name": "Время", "reporting_period_start": "2026-08-17",
        "fields": fields, "layout": {"data_rect": None, "header_rows": 1, "orientation": "columns", "skip_rows": []}})
    assert r.status_code == 201, r.text
    async with db.acquire() as conn:
        vals = await conn.fetch(
            "select v.row_label, v.value_number, v.value_text from dataset_values v "
            "join dataset_releases r on r.id=v.dataset_release_id "
            "where r.code='ztest_wait' and v.canonical_field_code=$1 order by v.row_label", wait["field_code"])
        tbl = await _widgetsources._dataset_table(conn, ids["org"], "ztest_wait")
    assert [round(float(v["value_number"]), 4) for v in vals] == [round(14 + 22 / 60, 4), round(87 + 16 / 60, 4)]
    assert vals[0]["value_text"] == "00:14:22", "исходная запись сохранена"
    assert tbl["column_units"] == {wait["field_code"]: "мин"}, "таблица знает, что графа — длительность"
