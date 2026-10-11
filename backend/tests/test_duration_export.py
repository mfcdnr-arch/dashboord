"""Длительности в выгрузке и в тексте — не дробью минут.

«87,267 мин» заказчик прочитал как 87 тысяч минут (11.10.2026). В Excel
значение уходит временем «ч:мм:сс» (его можно складывать и усреднять), прирост —
текстом со знаком (отрицательное время Excel рисует «#####»), а в тексте —
«1 ч 27 мин 16 с».
"""
from __future__ import annotations

from datetime import timedelta

from openpyxl import Workbook

from app.modules.dashboards._widgetexport import DUR_FMT, _dump_widget
from app.modules.ingestion.analyze import fmt_duration


def _book():
    wb = Workbook()
    summary = wb.active
    summary.append(["Виджет", "Тип", "Показатель", "Значение"])
    names = iter(f"Лист {i}" for i in range(10))
    return wb, summary, lambda _wid, _base: next(names)


def _as_td(cell) -> timedelta:
    assert cell.number_format == DUR_FMT, cell.number_format
    return timedelta(days=cell.value)


def test_fmt_duration_keeps_seconds_and_sign():
    assert fmt_duration(87.267) == "1 ч 27 мин 16 с"
    assert fmt_duration(14 + 22 / 60) == "14 мин 22 с"
    assert fmt_duration(0.75) == "45 с"
    assert fmt_duration(60) == "1 ч"
    assert fmt_duration(-3.7, signed=True) == "−3 мин 42 с"
    assert fmt_duration(3.7, signed=True) == "+3 мин 42 с"
    assert fmt_duration(None) == ""


def test_kpi_and_ranked_go_out_as_excel_time():
    wb, summary, sheet = _book()
    _dump_widget(wb, summary, sheet, "w1", "kpi", "Среднее ожидание",
                 {"type": "kpi", "value": 87.267, "unit": "мин"})
    assert _as_td(summary.cell(row=2, column=4)) == timedelta(hours=1, minutes=27, seconds=16)
    assert "ч:мм:сс" in summary.cell(row=2, column=3).value

    _dump_widget(wb, summary, sheet, "w2", "ranked", "Где дольше",
                 {"type": "ranked", "unit": None, "rows": [{"rank": 1, "label": "А", "value": 364.1}]},
                 unit="мин")  # единица только в настройке виджета — тоже работает
    ws = wb["Лист 0"]
    assert _as_td(ws.cell(row=2, column=3)) == timedelta(hours=6, minutes=4, seconds=6)
    assert ws.cell(row=2, column=1).value == 1  # место — не время


def test_dynamics_changes_are_signed_text_and_counts_stay_numbers():
    wb, summary, sheet = _book()
    _dump_widget(wb, summary, sheet, "w", "dynamics", "Динамика", {
        "type": "dynamics", "unit": "мин", "periods": ["2026-08-17", "2026-08-18"], "values": [17.42, 13.7],
        "total_change": -3.72, "total_change_pct": -21.4, "first_period": "2026-08-17",
        "last_period": "2026-08-18"})
    ws = wb["Лист 0"]
    assert _as_td(ws.cell(row=2, column=2)) == timedelta(minutes=17, seconds=25)
    assert ws.cell(row=5, column=2).value == "−3 мин 43 с"

    # Талоны — не минуты: число остаётся числом.
    _dump_widget(wb, summary, sheet, "k", "kpi", "Талоны", {"type": "kpi", "value": 10085, "unit": None})
    assert summary.cell(row=summary.max_row, column=4).value == 10085


def test_table_converts_only_columns_measured_in_minutes():
    wb, summary, sheet = _book()
    _dump_widget(wb, summary, sheet, "t", "table", "Таблица", {
        "type": "table", "columns": ["wait", "tickets"], "column_titles": {"wait": "Ожидание", "tickets": "Талоны"},
        "column_units": {"wait": "мин"}, "rows": [{"row": "А", "wait": 7, "tickets": 531}]})
    ws = wb["Лист 0"]
    assert _as_td(ws.cell(row=2, column=2)) == timedelta(minutes=7)
    assert ws.cell(row=2, column=3).value == 531
    assert ws.cell(row=1, column=2).value == "Ожидание (ч:мм:сс)"
    assert ws.cell(row=1, column=3).value == "Талоны"
